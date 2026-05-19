import msoffcrypto
from openpyxl import load_workbook
from io import BytesIO
import csv

excel_file = r"D:\INIA_CODE\CarlaSoftware\CopyTextura\39. F-65 Reporte de Resultados Textura Ver.04 SU 1144-1203 SU 2M 07-05.xlsx"
password = "12"
output_csv = "textura_B_G_values.csv"

# Decrypt Excel file
decrypted = BytesIO()

with open(excel_file, "rb") as f:
    office_file = msoffcrypto.OfficeFile(f)
    office_file.load_key(password=password)
    office_file.decrypt(decrypted)

decrypted.seek(0)

# Read decrypted workbook
wb = load_workbook(decrypted, data_only=True)
ws = wb.active

start_row = 9
data = []

row = start_row
while True:
    value_b = ws[f"B{row}"].value
    value_g = ws[f"G{row}"].value

    # Stop when column B is empty
    if value_b is None or str(value_b).strip() == "":
        break

    data.append([row, value_b, value_g])
    row += 1

# Save CSV
with open(output_csv, "w", newline="", encoding="utf-8-sig") as f:
    writer = csv.writer(f)
    writer.writerow(["row", "B", "G"])
    writer.writerows(data)

print(f"Saved: {output_csv}")

# Also print values
for row_num, value_b, value_g in data:
    print(f"Row {row_num}: B={value_b}, G={value_g}")
