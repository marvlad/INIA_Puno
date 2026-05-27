# extract_riqueza_fert.py

from pathlib import Path
from openpyxl import load_workbook
import csv


INPUT_EXCEL = "Software_Mejorado_Cultivos_Anuales_2025-2026_Arapa_CORREGIDO.xlsx"
OUTPUT_CSV = "Riqueza_Fert.csv"

POSSIBLE_SHEET_NAMES = [
    "Extracted_Fert",
]


def cell_value_to_text(value):
    if value is None:
        return ""

    return str(value).strip()


def find_sheet_name(wb):
    for name in POSSIBLE_SHEET_NAMES:
        if name in wb.sheetnames:
            return name

    print("\nAvailable sheets:")
    for sheet in wb.sheetnames:
        print(f"  - {sheet}")

    raise ValueError(
        "Could not find sheet. Expected one of: "
        + ", ".join(POSSIBLE_SHEET_NAMES)
    )


def extract_sheet_to_csv(
    input_excel=INPUT_EXCEL,
    output_csv=OUTPUT_CSV,
):
    input_excel = Path(input_excel)
    output_csv = Path(output_csv)

    if not input_excel.exists():
        raise FileNotFoundError(f"Excel file not found: {input_excel}")

    wb = None

    try:
        wb = load_workbook(
            input_excel,
            data_only=True,
            read_only=True,
        )

        sheet_name = find_sheet_name(wb)
        ws = wb[sheet_name]

        print(f"Reading Excel file: {input_excel}")
        print(f"Sheet found: {sheet_name}")

        rows = []

        for row in ws.iter_rows(values_only=True):
            values = [
                cell_value_to_text(value)
                for value in row
            ]

            if all(value == "" for value in values):
                continue

            rows.append(values)

        if not rows:
            raise ValueError(f"No data found in sheet: {sheet_name}")

        max_len = max(len(row) for row in rows)

        clean_rows = []
        for row in rows:
            row = row + [""] * (max_len - len(row))
            clean_rows.append(row)

        output_csv.parent.mkdir(parents=True, exist_ok=True)

        with open(output_csv, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerows(clean_rows)

        print("\nCSV created successfully.")
        print(f"Output CSV: {output_csv}")
        print(f"Rows written: {len(clean_rows)}")
        print(f"Columns written: {max_len}")

    finally:
        if wb is not None:
            wb.close()


if __name__ == "__main__":
    extract_sheet_to_csv()
